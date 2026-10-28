# -*- coding: utf-8 -*-
# Copyright 2026 Alain Bloos <alainbloos@gmail.com>
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl).

import base64
import datetime
import io
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer
from unittest.mock import patch

from odoo.tests import tagged
from odoo.tests.common import TransactionCase

from odoo.addons.os_ai_fields import patch as ai_patch
from odoo.addons.os_ai_fields.helpers import resolve_prompt


@tagged('post_install', '-at_install')
class TestValueConversion(TransactionCase):
    """LLM answers are validated and converted before touching any record."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Partner = cls.env['res.partner']
        cls.partner = cls.Partner.create({'name': 'Conversion partner'})

    def convert(self, fname, value, candidates=None):
        return ai_patch._convert_value(self.partner, self.Partner._fields[fname], value, candidates)

    def test_integer(self):
        self.assertEqual(self.convert('color', 4), 4)
        self.assertEqual(self.convert('color', '7'), 7)
        self.assertEqual(self.convert('color', 'Score: 9'), 9)
        self.assertEqual(self.convert('color', 3.0), 3)
        with self.assertRaises(ValueError):
            self.convert('color', 'no number here')
        with self.assertRaises(ValueError):
            self.convert('color', 'between 3 and 5')
        with self.assertRaises(ValueError):
            self.convert('color', 9000112310020011)  # seen from a small local model

    def test_float(self):
        self.assertEqual(self.convert('partner_latitude', 7.5), 7.5)
        self.assertEqual(self.convert('partner_latitude', '-12.25'), -12.25)
        with self.assertRaises(ValueError):
            self.convert('partner_latitude', 'unknown')

    def test_boolean(self):
        self.assertIs(self.convert('is_company', True), True)
        self.assertIs(self.convert('is_company', 'yes'), True)
        self.assertIs(self.convert('is_company', 'sí'), True)
        self.assertIs(self.convert('is_company', 'false'), False)
        with self.assertRaises(ValueError):
            self.convert('is_company', 'maybe')

    def test_selection_maps_to_canonical_key(self):
        self.assertEqual(self.convert('type', 'invoice'), 'invoice')
        self.assertEqual(self.convert('type', 'INVOICE'), 'invoice')
        label = dict(self.Partner._fields['type']._description_selection(self.env))['delivery']
        self.assertEqual(self.convert('type', label), 'delivery')
        with self.assertRaises(ValueError):
            self.convert('type', 'not-a-key')

    def test_explicit_null(self):
        self.assertIs(self.convert('name', None), False)
        self.assertIs(self.convert('type', None), False)
        self.assertIs(self.convert('type', ''), False)  # "none" sentinel of the selection enum
        self.assertIs(self.convert('country_id', '0', {self.env.ref('base.ar').id}), False)

    def test_many2one_against_candidates(self):
        ar = self.env.ref('base.ar')
        fr = self.env.ref('base.fr')
        self.assertEqual(self.convert('country_id', ar.id, {ar.id}), ar.id)
        self.assertEqual(self.convert('country_id', str(ar.id), {ar.id}), ar.id)
        self.assertIs(self.convert('country_id', 0, {ar.id}), False)
        with self.assertRaises(ValueError):
            self.convert('country_id', fr.id, {ar.id})
        with self.assertRaises(ValueError):
            self.convert('country_id', ar.id, set())

    def test_many2many_against_candidates(self):
        tags = self.env['res.partner.category'].create([{'name': 'AI tag %s' % i} for i in range(3)])
        candidates = set(tags.ids)
        self.assertEqual(self.convert('category_id', [tags[0].id, str(tags[1].id), tags[0].id], candidates),
                         [tags[0].id, tags[1].id])
        self.assertEqual(self.convert('category_id', '%s, %s' % (tags[2].id, tags[0].id), candidates),
                         [tags[2].id, tags[0].id])
        self.assertEqual(self.convert('category_id', 'NONE', candidates), [])
        self.assertEqual(self.convert('category_id', [], candidates), [])
        with self.assertRaises(ValueError):
            self.convert('category_id', [tags[0].id, 999999], candidates)

    def test_date(self):
        date_field = self.env['res.currency.rate']._fields['name']
        convert = ai_patch._convert_value
        self.assertEqual(convert(self.partner, date_field, '2026-10-03'), datetime.date(2026, 10, 3))
        self.assertEqual(convert(self.partner, date_field, '2026-10-03T00:00:00'), datetime.date(2026, 10, 3))
        for value in ('next tuesday', '2026-10-03 and more', '2026-10-03T15:30:00'):
            with self.assertRaises(ValueError, msg=value):
                convert(self.partner, date_field, value)

    def test_datetime_offsets_are_converted_to_utc(self):
        """An offset is not dropped: the instant is kept."""
        field = self.Partner._fields['write_date']
        convert = ai_patch._convert_value
        expected = datetime.datetime(2026, 10, 24, 9, 0, 0)
        for value in ('2026-10-24T12:00:00+03:00', '2026-10-24T09:00:00Z', '2026-10-24 06:00:00-03:00',
                      '2026-10-24T09:00:00.250000+00:00', '2026-10-24 09:00:00'):
            self.assertEqual(convert(self.partner, field, value), expected, value)
        for value in ('2026-10-24T09:00:00 UTC', 'tomorrow', '2026-10-24T25:00:00'):
            with self.assertRaises(ValueError, msg=value):
                convert(self.partner, field, value)

    def test_numbers_without_leading_zero(self):
        """.5 is 0.5, not 5."""
        for value, expected in (('.5', 0.5), ('-.5', -0.5), ('approximately .5', 0.5), ('+2.5', 2.5),
                                ('1e3', 1000.0), ('-12.25', -12.25)):
            self.assertEqual(self.convert('partner_latitude', value), expected, value)

    def test_char(self):
        self.assertEqual(self.convert('name', 'Hello'), 'Hello')
        self.assertEqual(self.convert('name', 42), '42')
        with self.assertRaises(ValueError):
            self.convert('name', {'nested': 'object'})


@tagged('post_install', '-at_install')
class TestPromptPlaceholders(TransactionCase):

    def test_company_placeholder_uses_record_company(self):
        other_company = self.env['res.company'].create({'name': 'Placeholder Second Co'})
        partner = self.env['res.partner'].create({'name': 'Company partner', 'company_id': other_company.id})
        text, _images = resolve_prompt('Company: {__company__}', partner)
        self.assertEqual(text, 'Company: Placeholder Second Co')

    def test_company_placeholder_falls_back_to_env_company(self):
        partner = self.env['res.partner'].create({'name': 'Shared partner', 'company_id': False})
        text, _images = resolve_prompt('{__company__}', partner)
        self.assertEqual(text, self.env.company.name)


@tagged('post_install', '-at_install')
class TestValueSchemas(TransactionCase):
    """Closed sets are requested as enums."""

    def schema(self, fname, candidates=None):
        return ai_patch._value_schema(self.env, self.env['res.partner']._fields[fname], candidates)

    def test_selection_enum(self):
        keys = [key for key, _label in self.env['res.partner']._fields['type']._description_selection(self.env)]
        self.assertEqual(self.schema('type'), {'type': 'string', 'enum': keys + ['']})

    def test_many2one_enum(self):
        self.assertEqual(self.schema('country_id', {12, 3}), {'type': 'string', 'enum': ['0', '3', '12']})

    def test_many2many_enum(self):
        self.assertEqual(self.schema('category_id', {7, 5}),
                         {'type': 'array', 'items': {'type': 'string', 'enum': ['5', '7']}})

    def test_scalar_types(self):
        self.assertEqual(self.schema('color'), {'type': ['integer', 'null']})
        self.assertEqual(self.schema('is_company'), {'type': ['boolean', 'null']})
        self.assertEqual(self.schema('name'), {'type': ['string', 'null']})


@tagged('post_install', '-at_install')
class TestDocuments(TransactionCase):
    """Documents are produced with the report engine available in each Odoo version."""

    def test_html_to_pdf(self):
        from odoo.addons.os_ai_fields import document_utils
        pdf_b64 = document_utils.html_to_pdf(self.env, '<h1>AI report</h1><p>Generated in a test.</p>')
        self.assertTrue(pdf_b64, 'no PDF produced (is wkhtmltopdf installed?)')
        import base64
        self.assertTrue(base64.b64decode(pdf_b64).startswith(b'%PDF'))

    def test_json_to_xlsx(self):
        from odoo.addons.os_ai_fields import document_utils
        xlsx_b64 = document_utils.json_to_xlsx('{"headers": ["Field", "Value"], "rows": [["Name", "ACME"], ["Score", "42"]]}')
        self.assertTrue(xlsx_b64)
        import base64
        self.assertTrue(base64.b64decode(xlsx_b64).startswith(b'PK'))  # xlsx is a zip file


@tagged('post_install', '-at_install')
class TestAnswerParsing(TransactionCase):
    """Answers that cannot be trusted never become values."""

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        cls.Partner = cls.env['res.partner']
        cls.partner = cls.Partner.create({'name': 'Parsing partner'})

    def extract(self, fname, content, candidates=None, multi_lang=False, lang_codes=()):
        return ai_patch._extract_single_value(
            self.partner, self.Partner._fields[fname], content, 'capable_text', None,
            multi_lang, list(lang_codes), candidates)

    def test_truncated_json_in_a_fence_is_rejected(self):
        for content in ('```json\n{"value":"cut\n```', '```json\n{"value":"cut', '```\n{"value":"cut\n```',
                        '  ```JSON \n{"value": "cut"', '{"value":"cut'):
            with self.assertRaises(ValueError, msg=content):
                self.extract('comment', content)

    def test_fenced_answers(self):
        self.assertEqual(self.extract('comment', '```json\n{"value": "ok"}\n```'), 'ok')
        self.assertEqual(self.extract('comment', '```\nplain text\n```'), 'plain text')

    def test_duplicated_keys_are_rejected(self):
        with self.assertRaises(ValueError):
            self.extract('comment', '{"value": "A", "value": "B"}')
        with self.assertRaises(ValueError):
            ai_patch._parse_batch_answers('{"12": "A", "12": "B"}', [12], False, [])

    def test_duplicated_languages_are_rejected(self):
        with self.assertRaises(ValueError):
            ai_patch._translations_from({'translations': [
                {'lang_code': 'en_US', 'value': 'A'}, {'lang_code': 'en_US', 'value': 'B'}]})
        answers = ai_patch._parse_batch_answers(
            '{"results": [{"id": 12, "translations": [{"lang_code": "en_US", "value": "A"}, '
            '{"lang_code": "en_US", "value": "B"}]}]}', [12], True, ['en_US'])
        self.assertIs(answers['12'], ai_patch._DUPLICATE)

    def test_many2many_answer_as_a_list(self):
        tags = self.env['res.partner.category'].create([{'name': 'Parsing tag %s' % i} for i in range(2)])
        self.assertEqual(self.extract('category_id', '["%s", "%s"]' % tuple(tags.ids), set(tags.ids)), tags.ids)
        with self.assertRaises(ValueError):
            self.extract('name', '["a", "b"]')


@tagged('post_install', '-at_install')
class TestDocumentSafety(TransactionCase):
    """Generated documents are untrusted input."""

    def setUp(self):
        super().setUp()
        from odoo.addons.os_ai_fields import document_utils
        self.docs = document_utils

    def test_html_without_remote_resources(self):
        image = 'data:image/png;base64,iVBORw0KGgo='
        html = self.docs.sanitize_html(
            '<html><head><link rel="stylesheet" href="http://127.0.0.1:9/x.css"><base href="http://h/">'
            '<style>@import url(http://h/a.css); body {background: url("http://h/b.png"); color: red}</style>'
            '</head><body background="http://h/c.png" onload="x()"><h1 style="background:url(http://h/d.png)">T</h1>'
            '<img src="http://127.0.0.1:9/e.png"><img src="%s" alt="kept"><img srcset="http://h/f 2x" src="%s">'
            '<iframe src="file:///etc/passwd"></iframe><script>fetch("http://h")</script><object data="http://h/o">'
            '</object><svg><image href="http://h/s.png"/></svg><a href="javascript:alert(1)">js</a>'
            '<a href="https://example.com">link</a><table><tr><td style="border: 1px solid #ccc">cell</td></tr>'
            '</table></body></html>' % (image, image))
        for forbidden in ('http://h', '127.0.0.1', 'file:', 'javascript:', '<script', '<iframe', '<object',
                          '<svg', '<link', '<base', 'onload', 'srcset', '@import'):
            self.assertNotIn(forbidden, html)
        self.assertIn('color: red', html)
        self.assertIn(image, html)
        self.assertIn('href="https://example.com"', html)
        self.assertIn('<td style="border: 1px solid #ccc">cell</td>', html)
        self.assertIn('<meta charset="utf-8">', html)

    def test_pdf_engine_fetches_nothing(self):
        """Real report engine: the generated HTML cannot make the server fetch a URL."""

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def do_GET(self):
                self.server.paths.append(self.path)
                self.send_response(404)
                self.end_headers()

        server = HTTPServer(('127.0.0.1', 0), Handler)
        server.paths = []
        threading.Thread(target=server.serve_forever, daemon=True).start()
        try:
            url = 'http://127.0.0.1:%d' % server.server_address[1]
            pdf = self.docs.html_to_pdf(self.env, (
                '<h1>Report</h1><img src="%s/img.png"><div style="background: url(%s/bg.png)">x</div>'
                '<link rel="stylesheet" href="%s/style.css">' % (url, url, url)))
        finally:
            server.shutdown()
            server.server_close()
        self.assertTrue(pdf and base64.b64decode(pdf).startswith(b'%PDF'), 'no PDF (is wkhtmltopdf installed?)')
        self.assertEqual(server.paths, [])

    def test_xlsx_cells_are_never_formulas(self):
        from openpyxl import load_workbook
        xlsx = self.docs.json_to_xlsx(
            '{"headers": ["=SUM(A1)", "Phone"], "rows": [["=1+1", "+5491112345678"], ["00123", "42"], '
            '["5491112345678", "3.5"], [true, null], [], ["=2+2"]]}')
        sheet = load_workbook(io.BytesIO(base64.b64decode(xlsx))).active
        cells = {cell.coordinate: (cell.value, cell.data_type) for row in sheet.iter_rows() for cell in row}
        self.assertEqual(cells['A1'], ('=SUM(A1)', 's'))
        self.assertEqual(cells['A2'], ('=1+1', 's'))
        self.assertEqual(cells['B2'], ('+5491112345678', 's'))
        self.assertEqual(cells['A3'], ('00123', 's'))
        self.assertEqual(cells['B3'], (42, 'n'))
        self.assertEqual(cells['A4'], ('5491112345678', 's'))  # long references stay text
        self.assertEqual(cells['B4'], (3.5, 'n'))
        self.assertEqual(cells['A5'][0], True)
        self.assertEqual(cells['A7'], ('=2+2', 's'))  # also after an empty row

    def test_xlsx_structure_is_validated(self):
        for content in ('{"headers": "Name", "rows": []}', '{"headers": ["A"], "rows": [[{"nested": 1}]]}',
                        '{"headers": ["A"], "rows": "x"}',
                        '{"headers": ["A"], "rows": [["FIRST"]], "rows": [["LAST"]]}'):
            self.assertFalse(self.docs.json_to_xlsx(content), content)

    def test_pdf_answer_contract(self):
        env = self.env
        with patch.object(self.docs, 'html_to_pdf', return_value='UERG') as html_to_pdf:
            self.assertEqual(ai_patch._postprocess_document(env, 'pdf', '{"value": "<h1>Hi</h1>"}'), 'UERG')
            self.assertEqual(html_to_pdf.call_args.args[1], '<h1>Hi</h1>')
            self.assertEqual(ai_patch._postprocess_document(env, 'pdf', '<h1>Plain HTML</h1>'), 'UERG')
            for content in ('{"unexpected": "payload"}', '{"value": 42}', '[1, 2]', '{"value": "cut'):
                with self.assertRaises(ValueError, msg=content):
                    ai_patch._postprocess_document(env, 'pdf', content)


@tagged('post_install', '-at_install')
class TestPromptResolution(TransactionCase):
    """Placeholders resolve to the record data, or fail explicitly."""

    def test_values(self):
        parent = self.env['res.partner'].create({'name': 'Resolution parent', 'is_company': True})
        self.env['res.partner'].create([{'name': 'Child A', 'parent_id': parent.id}, {'name': 'Child B', 'parent_id': parent.id}])
        orphan = self.env['res.partner'].create({'name': 'Orphan'})
        text, _images = resolve_prompt('{name} / {parent_id.name} / {child_ids.name}', parent)
        self.assertEqual(text, 'Resolution parent /  / Child A, Child B')
        text, _images = resolve_prompt('Parent: {parent_id.name}. Return {"value": "x"}.', orphan)
        self.assertEqual(text, 'Parent: . Return {"value": "x"}.')  # empty relation; JSON stays literal

    def test_zeros_and_booleans_are_values(self):
        """0 and False are values in the prompt, not missing data."""
        parent = self.env['res.partner'].create({'name': 'Coordinates parent', 'is_company': True})
        self.env['res.partner'].create([
            {'name': 'Zero', 'parent_id': parent.id, 'partner_latitude': 0.0},
            {'name': 'Other', 'parent_id': parent.id, 'partner_latitude': 2.5},
        ])
        text, _images = resolve_prompt('{child_ids.partner_latitude}', parent)
        self.assertCountEqual(text.split(', '), ['0.0', '2.5'])
        person = self.env['res.partner'].create({'name': 'Person', 'is_company': False})
        self.assertEqual(resolve_prompt('Company: {is_company}', person)[0], 'Company: False')
        self.assertEqual(resolve_prompt('{child_ids.is_company}', parent)[0], 'False, False')

    def test_errors_are_explicit(self):
        partner = self.env['res.partner'].create({'name': 'Resolution errors'})
        for template in ('{nmae}', '{name.first}', '{parent_id.nmae}'):
            with self.assertRaises(ValueError, msg=template):
                resolve_prompt(template, partner)

    def test_field_configuration_is_checked(self):
        Partner = self.env['res.partner']
        field = Partner._fields['ai_greeting'] if 'ai_greeting' in Partner._fields else Partner._fields['comment']
        self.assertIsNone(ai_patch._check_placeholders(self.env, 'res.partner', '{name} {parent_id.country_id.code}'))
        self.assertIn('nmae', ai_patch._check_placeholders(self.env, 'res.partner', 'Hi {nmae}'))
        self.assertIn('not a relational field', ai_patch._check_placeholders(self.env, 'res.partner', '{name.x}'))
        self.assertIsNone(ai_patch._check_placeholders(self.env, 'res.partner', 'JSON {"value": 1} {__user__}'))
        with patch.object(field, 'ai_compute', 'Hello {name}'), patch.object(field, 'compute_sudo', True), \
                patch.object(field, 'ai_compute_engine', 'decison'):
            plan, reason = ai_patch.prepare_computation(self.env, 'res.partner', field.name, field, 1)
        self.assertIsNone(plan)
        self.assertIn('unknown ai_compute_engine', reason)

    def test_uncovered_placeholders(self):
        """Dotted paths must be listed; relational placeholders need a path through them."""
        self.assertEqual(ai_patch._uncovered_placeholders(
            self.env, 'res.partner', '{name} {parent_id.name} {country_id.code} {category_id} {child_ids}',
            ['name', 'parent_id', 'country_id.code', 'child_ids.name']),
            ['category_id (display names)', 'parent_id.name'])


@tagged('post_install', '-at_install')
class TestDecisionRequest(TransactionCase):
    """Decision requests keep every placeholder apart."""

    def test_distinct_placeholders_get_distinct_keys(self):
        partner = self.env['res.partner'].create({'name': 'Decision keys'})
        field = self.env['res.partner']._fields['is_company']
        with patch('odoo.addons.os_ai_fields.helpers.resolve_prompt', side_effect=lambda text, record: (text, [])):
            state, questions, _options = ai_patch._decision_request(
                partner, field, 'Compare {partner_id.name} with {partner_id_name} and {partner_id.name}')
        self.assertEqual(state, {'partner_id_name': '{partner_id.name}', 'partner_id_name_2': '{partner_id_name}'})
        self.assertEqual(questions['value']['instructions'],
                         'Compare `partner_id_name` with `partner_id_name_2` and `partner_id_name`')

    def test_reserved_selection_key(self):
        partner = self.env['res.partner'].create({'name': 'Reserved key'})
        field = self.env['res.partner']._fields['type']
        with patch.object(type(field), '_description_selection',
                          lambda self, env: [('a', 'A'), (ai_patch.DECISION_NONE, 'Clash')]):
            with self.assertRaises(ValueError):
                ai_patch._decision_request(partner, field, 'Which type is {name}?')
