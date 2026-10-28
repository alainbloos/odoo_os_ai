# -*- coding: utf-8 -*-
# Copyright 2026 Alain Bloos <alainbloos@gmail.com>
# License LGPL-3.0 or later (https://www.gnu.org/licenses/lgpl).

"""
Post-processing utilities for AI-generated documents.

These functions convert structured LLM output (JSON / HTML) into binary
documents with openpyxl (declared in the manifest) and the database's report
engine. The model output is untrusted input: spreadsheet cells never become
formulas and the HTML is sanitized before the PDF engine sees it.

  json_to_xlsx  — JSON {"headers": [...], "rows": [[...], ...]} → base64 xlsx
  html_to_pdf   — HTML string → base64 pdf  (report engine of the database)
"""

import base64
import io
import json
import logging
import math
import re

from lxml import html as lxml_html

from .compat import render_pdf

_logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Excel  (openpyxl)
# ---------------------------------------------------------------------------

def json_to_xlsx(content_str):
    """Convert a JSON string with *headers* + *rows* into a base64 xlsx.

    Expected JSON structure::

        {
            "headers": ["Name", "Email", "Score"],
            "rows": [
                ["Alice", "alice@example.com", 85],
                ["Bob",   "bob@example.com",   72]
            ]
        }

    Returns:
        str: base64-encoded xlsx bytes, or ``False`` on error.
    """
    try:
        from openpyxl import Workbook
    except ImportError:
        _logger.error("os_ai_fields: openpyxl is not installed — cannot generate xlsx")
        return False

    # -- Parse JSON --------------------------------------------------------
    try:
        cleaned = _strip_markdown_fences(content_str)
        data = json.loads(cleaned, object_pairs_hook=_reject_duplicate_keys)
    except (TypeError, ValueError) as exc:
        _logger.error("os_ai_fields: json_to_xlsx — invalid JSON: %s\n%s",
                       exc, (content_str or '')[:500])
        return False

    if not isinstance(data, dict):
        _logger.error("os_ai_fields: json_to_xlsx — expected dict, got %s", type(data).__name__)
        return False

    headers = data.get('headers') or []
    rows = data.get('rows') or []
    if not isinstance(headers, list) or not isinstance(rows, list):
        _logger.error("os_ai_fields: json_to_xlsx — headers and rows must be lists")
        return False
    rows = [row if isinstance(row, list) else [row] for row in rows]
    if not all(_is_cell(value) for value in headers) or not all(_is_cell(value) for row in rows for value in row):
        _logger.error("os_ai_fields: json_to_xlsx — cells must be text, numbers, booleans or null")
        return False

    if not headers and not rows:
        _logger.warning("os_ai_fields: json_to_xlsx — empty headers and rows")
        return False

    # -- Build workbook ----------------------------------------------------
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    if headers:
        _append_row(ws, ['' if h is None else str(h) for h in headers])
        # Bold header row
        from openpyxl.styles import Font
        bold = Font(bold=True)
        for cell in ws[1]:
            cell.font = bold

    for row in rows:
        _append_row(ws, [_coerce_cell_value(v) for v in row])

    # Auto-fit column widths (approximate)
    for col in ws.columns:
        max_len = 0
        col_letter = col[0].column_letter
        for cell in col:
            val = str(cell.value) if cell.value is not None else ''
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = min(max_len + 3, 50)

    # -- Serialize to bytes ------------------------------------------------
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return base64.b64encode(buf.read()).decode('ascii')


# ---------------------------------------------------------------------------
# PDF  (report engine of the database: wkhtmltopdf, Paper Muncher in 20+...)
# ---------------------------------------------------------------------------

def html_to_pdf(env, html_str):
    """Convert an HTML string into a base64-encoded PDF.

    Uses the report engine of the database (wkhtmltopdf, or the engine
    configured in Odoo 20+), see ``compat.render_pdf``.

    Args:
        env: Odoo Environment (used to access ``ir.actions.report``).
        html_str: HTML content string.

    Returns:
        str: base64-encoded PDF bytes, or ``False`` on error.
    """
    if not html_str or not html_str.strip():
        _logger.warning("os_ai_fields: html_to_pdf — empty HTML input")
        return False

    cleaned = _strip_markdown_fences(html_str)

    # Wrap in a full HTML document if the LLM returned a fragment
    if '<html' not in cleaned.lower():
        cleaned = _wrap_html_fragment(cleaned)

    try:
        cleaned = sanitize_html(cleaned)
    except Exception as exc:
        _logger.error("os_ai_fields: html_to_pdf — the HTML cannot be parsed: %s", exc)
        return False

    try:
        pdf_bytes = render_pdf(env, cleaned)
        return base64.b64encode(pdf_bytes).decode('ascii')
    except Exception as exc:
        _logger.error("os_ai_fields: html_to_pdf — PDF rendering failed: %s", exc)
        return False


# ---------------------------------------------------------------------------
# Internal helpers
# ---------------------------------------------------------------------------

# A plain decimal number: no sign other than "-", no leading zeros, no exponent
# and at most 11 integer digits. Anything else (00123, +5491112345678, long
# references) is an identifier and stays text.
_CELL_NUMBER_RE = re.compile(r'-?(?:0|[1-9]\d{0,10})(?:\.\d{1,10})?')


def _reject_duplicate_keys(pairs):
    """json.loads keeps the last of repeated keys: refuse them instead."""
    result = {}
    for key, value in pairs:
        if key in result:
            raise ValueError('the answer repeats the key %r' % key)
        result[key] = value
    return result


def _is_cell(value):
    return value is None or isinstance(value, (str, bool, int)) or (isinstance(value, float) and math.isfinite(value))


def _coerce_cell_value(value):
    """Spreadsheet value of a cell: plain decimal numbers and true/false are converted."""
    if not isinstance(value, str):
        return value
    text = value.strip()
    if _CELL_NUMBER_RE.fullmatch(text):
        return float(text) if '.' in text else int(text)
    if text.lower() in ('true', 'false'):
        return text.lower() == 'true'
    return value


def _append_row(ws, values):
    """Append a row; text is always stored as text, never as a formula."""
    ws.append(values)
    for cell, value in zip(ws[ws.max_row], values):
        if isinstance(value, str):
            cell.data_type = 's'


# ---------------------------------------------------------------------------
# HTML sanitizer for the PDF engine
# ---------------------------------------------------------------------------

# Elements removed with their content: scripts, embedded or remote content,
# forms and vector graphics (SVG can reference external resources).
_DROPPED_TAGS = frozenset({
    'script', 'iframe', 'frame', 'frameset', 'object', 'embed', 'applet', 'link', 'base', 'meta',
    'svg', 'math', 'audio', 'video', 'source', 'track', 'picture', 'canvas', 'template', 'noscript',
    'input', 'button', 'select', 'textarea', 'portal', 'param',
})
# Elements kept; other tags are removed but their text is kept.
_SAFE_TAGS = frozenset({
    'html', 'head', 'body', 'title', 'style', 'div', 'span', 'p', 'br', 'hr', 'h1', 'h2', 'h3', 'h4',
    'h5', 'h6', 'b', 'strong', 'i', 'em', 'u', 's', 'strike', 'del', 'ins', 'small', 'big', 'sub', 'sup',
    'mark', 'abbr', 'cite', 'q', 'blockquote', 'pre', 'code', 'kbd', 'samp', 'var', 'address', 'time',
    'ul', 'ol', 'li', 'dl', 'dt', 'dd', 'table', 'thead', 'tbody', 'tfoot', 'tr', 'th', 'td', 'caption',
    'colgroup', 'col', 'a', 'img', 'section', 'article', 'header', 'footer', 'main', 'nav', 'aside',
    'figure', 'figcaption', 'center', 'font', 'label', 'details', 'summary',
})
_SAFE_ATTRS = frozenset({
    'class', 'id', 'style', 'title', 'alt', 'lang', 'dir', 'colspan', 'rowspan', 'align', 'valign',
    'width', 'height', 'border', 'cellpadding', 'cellspacing', 'scope', 'headers', 'start', 'type',
    'reversed', 'span', 'href', 'src', 'name', 'abbr', 'datetime', 'summary', 'bgcolor', 'color',
    'face', 'size', 'media',
})
_DATA_IMAGE_RE = re.compile(r'data:image/(?:png|jpe?g|gif|webp);base64,[A-Za-z0-9+/=\s]+', re.IGNORECASE)
_SAFE_HREF_RE = re.compile(r'(?:https?:|mailto:|#)', re.IGNORECASE)
_CSS_COMMENT_RE = re.compile(r'/\*.*?\*/', re.DOTALL)
_CSS_IMPORT_RE = re.compile(r'@import[^;]*;?', re.IGNORECASE)
_CSS_IMAGE_SET_RE = re.compile(r'(?:-webkit-)?image-set\s*\([^)]*\)', re.IGNORECASE)
_CSS_URL_RE = re.compile(r'url\s*\(\s*([\'"]?)(.*?)\1\s*\)', re.IGNORECASE | re.DOTALL)


def _sanitize_css(css):
    """CSS without remote resources: url() other than embedded images, @import, image-set."""
    css = _CSS_COMMENT_RE.sub('', css or '')
    if '\\' in css:
        return ''  # escapes could hide url( or @import
    css = _CSS_IMPORT_RE.sub('', css)
    css = _CSS_IMAGE_SET_RE.sub('none', css)
    return _CSS_URL_RE.sub(
        lambda match: match.group(0) if _DATA_IMAGE_RE.fullmatch(match.group(2).strip()) else 'none', css)


def sanitize_html(document):
    """Make generated HTML safe for the PDF engine.

    The engine fetches every resource the HTML references, from the Odoo
    server: generated HTML must not reach other hosts, local services or
    files. Only formatting elements are kept, images must be embedded
    (data:image/... base64), CSS cannot load anything, and there are no
    scripts, frames, forms or event handlers. Links stay clickable: the
    engine does not fetch them.

    Returns:
        str: a complete HTML document (UTF-8).
    """
    doc = lxml_html.document_fromstring(document)
    for element in list(doc.iter()):
        if element.getparent() is None and element is not doc:
            continue  # already removed with an ancestor
        tag = element.tag.lower() if isinstance(element.tag, str) else None
        if tag is None:  # comment or processing instruction
            element.drop_tree()
            continue
        if tag in _DROPPED_TAGS:
            element.drop_tree()
            continue
        if tag not in _SAFE_TAGS:
            element.drop_tag()
            continue
        if tag == 'style':
            element.text = _sanitize_css(element.text_content())
            for child in list(element):
                element.remove(child)
        for name in list(element.attrib):
            value = element.attrib[name]
            lname = name.lower()
            if lname not in _SAFE_ATTRS:
                del element.attrib[name]
            elif lname == 'style':
                element.attrib[name] = _sanitize_css(value)
            elif lname == 'href' and not (tag == 'a' and _SAFE_HREF_RE.match(value.strip())):
                del element.attrib[name]
            elif lname == 'src' and not (tag == 'img' and _DATA_IMAGE_RE.fullmatch(value.strip())):
                del element.attrib[name]
        if tag == 'img' and 'src' not in element.attrib:
            element.drop_tree()

    head = doc.find('head')
    if head is None:
        head = lxml_html.Element('head')
        doc.insert(0, head)
    head.insert(0, lxml_html.Element('meta', charset='utf-8'))
    return lxml_html.tostring(doc, encoding='unicode', method='html', doctype='<!DOCTYPE html>')


def _strip_markdown_fences(text):
    """Remove markdown code fences (```json ... ```) if present."""
    if not text:
        return text
    text = text.strip()
    text = re.sub(r'^```(?:json|html|xml)?\s*\n?', '', text, flags=re.IGNORECASE)
    text = re.sub(r'\n?\s*```$', '', text)
    return text.strip()


def _wrap_html_fragment(fragment):
    """Wrap an HTML fragment in a minimal HTML document for wkhtmltopdf."""
    return (
        '<!DOCTYPE html>\n'
        '<html>\n'
        '<head>\n'
        '  <meta charset="utf-8"/>\n'
        '  <style>\n'
        '    body { font-family: Arial, Helvetica, sans-serif; font-size: 12px; margin: 20px; }\n'
        '    table { border-collapse: collapse; width: 100%%; }\n'
        '    th, td { border: 1px solid #ccc; padding: 6px 10px; text-align: left; }\n'
        '    th { background-color: #f5f5f5; font-weight: bold; }\n'
        '  </style>\n'
        '</head>\n'
        '<body>\n%s\n</body>\n'
        '</html>'
    ) % fragment
